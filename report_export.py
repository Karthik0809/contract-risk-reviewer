"""Render a RiskReport to a branded, hand-off-ready PDF or DOCX (returns bytes).

Used by the Streamlit UI download buttons. Both exporters degrade gracefully: if
reportlab / python-docx aren't installed, the functions raise a clear RuntimeError
the caller can surface instead of crashing.
"""

from __future__ import annotations

import io
from datetime import datetime

from contract_schema import RiskReport, Severity

try:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        HRFlowable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
    )
    REPORTLAB_AVAILABLE = True
except Exception:  # pragma: no cover
    REPORTLAB_AVAILABLE = False

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    DOCX_AVAILABLE = True
except Exception:  # pragma: no cover
    DOCX_AVAILABLE = False


_SEVERITY_HEX = {
    Severity.HIGH: "#d64545",
    Severity.MEDIUM: "#e08e0b",
    Severity.LOW: "#c9a227",
    Severity.INFO: "#4a7fb5",
}


def _esc(text: str) -> str:
    return (text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def report_to_pdf(report: RiskReport) -> bytes:
    if not REPORTLAB_AVAILABLE:
        raise RuntimeError("PDF export requires reportlab (pip install reportlab).")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=letter,
        leftMargin=0.8 * inch, rightMargin=0.8 * inch,
        topMargin=0.8 * inch, bottomMargin=0.8 * inch,
        title=f"Risk Report — {report.contract_name}",
    )
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=20, spaceAfter=4)
    meta = ParagraphStyle("meta", parent=styles["Normal"], fontSize=9, textColor=colors.grey)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=13, spaceBefore=14, spaceAfter=6)
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=10, leading=14, alignment=TA_LEFT)
    quote = ParagraphStyle("quote", parent=body, leftIndent=12, textColor=colors.HexColor("#444444"),
                           fontName="Helvetica-Oblique", spaceBefore=4, spaceAfter=4)

    flow = []
    flow.append(Paragraph("Contract Risk Report", h1))
    flow.append(Paragraph(
        f"{_esc(report.contract_name)} &nbsp;•&nbsp; generated {datetime.now():%Y-%m-%d %H:%M} "
        f"&nbsp;•&nbsp; {report.engine} engine", meta))
    flow.append(Spacer(1, 8))
    flow.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#dddddd")))

    ov = report.overall_risk
    flow.append(Paragraph(
        f'Overall risk: <font color="{_SEVERITY_HEX.get(ov, "#777")}"><b>{ov.value}</b></font> '
        f'&nbsp; ({len(report.findings)} findings · {report.high_count()} high · '
        f'{report.medium_count()} medium · {report.clause_count} clauses)', h2))

    flow.append(Paragraph("Executive summary", h2))
    flow.append(Paragraph(_esc(report.summary), body))

    flow.append(Paragraph("Flagged clauses", h2))
    if not report.findings:
        flow.append(Paragraph("No risky or non-standard clauses detected.", body))
    for i, f in enumerate(report.sorted_findings(), start=1):
        color = _SEVERITY_HEX.get(f.severity, "#777")
        loc = f.locator()
        loc_str = f" &nbsp; <font color='#888'>[{_esc(loc)}]</font>" if loc else ""
        flow.append(Paragraph(
            f'{i}. <font color="{color}"><b>{f.severity.value}</b></font> — '
            f'<b>{_esc(f.category)}</b>: {_esc(f.title)}{loc_str}', body))
        flow.append(Paragraph(f"<b>Why it's risky:</b> {_esc(f.rationale)}", body))
        if f.recommendation:
            flow.append(Paragraph(f"<b>Recommended action:</b> {_esc(f.recommendation)}", body))
        flow.append(Paragraph(f'“{_esc(f.quote)}”', quote))
        flow.append(Spacer(1, 6))

    doc.build(flow)
    return buf.getvalue()


def report_to_docx(report: RiskReport) -> bytes:
    if not DOCX_AVAILABLE:
        raise RuntimeError("DOCX export requires python-docx (pip install python-docx).")

    docx = Document()
    docx.add_heading("Contract Risk Report", level=0)
    p = docx.add_paragraph()
    run = p.add_run(f"{report.contract_name}  •  generated {datetime.now():%Y-%m-%d %H:%M}  •  {report.engine} engine")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    ov = report.overall_risk
    p = docx.add_paragraph()
    p.add_run("Overall risk: ").bold = True
    r = p.add_run(ov.value)
    r.bold = True
    hexv = _SEVERITY_HEX.get(ov, "#777777").lstrip("#")
    r.font.color.rgb = RGBColor(int(hexv[0:2], 16), int(hexv[2:4], 16), int(hexv[4:6], 16))
    p.add_run(
        f"   ({len(report.findings)} findings · {report.high_count()} high · "
        f"{report.medium_count()} medium · {report.clause_count} clauses)")

    docx.add_heading("Executive summary", level=1)
    docx.add_paragraph(report.summary)

    docx.add_heading("Flagged clauses", level=1)
    if not report.findings:
        docx.add_paragraph("No risky or non-standard clauses detected.")
    for i, f in enumerate(report.sorted_findings(), start=1):
        p = docx.add_paragraph()
        sev = p.add_run(f"{i}. [{f.severity.value}] ")
        sev.bold = True
        hexv = _SEVERITY_HEX.get(f.severity, "#777777").lstrip("#")
        sev.font.color.rgb = RGBColor(int(hexv[0:2], 16), int(hexv[2:4], 16), int(hexv[4:6], 16))
        cat = p.add_run(f"{f.category}: {f.title}")
        cat.bold = True
        loc = f.locator()
        if loc:
            p.add_run(f"  [{loc}]").italic = True
        docx.add_paragraph(f"Why it's risky: {f.rationale}")
        if f.recommendation:
            docx.add_paragraph(f"Recommended action: {f.recommendation}")
        q = docx.add_paragraph()
        qr = q.add_run(f"“{f.quote}”")
        qr.italic = True

    buf = io.BytesIO()
    docx.save(buf)
    return buf.getvalue()
